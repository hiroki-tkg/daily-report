"""OLIVE+ の上場一覧（画面コピペのテキスト）を読み取り、明細と集計の Excel を作る。

使い方:
    python parse_offer.py samples/*.txt -o out.xlsx
"""
import argparse
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PREFECTURES = [
    "北海道", "青森", "岩手", "宮城", "秋田", "山形", "福島", "茨城", "栃木", "群馬",
    "埼玉", "千葉", "東京", "神奈川", "新潟", "富山", "石川", "福井", "山梨", "長野",
    "岐阜", "静岡", "愛知", "三重", "滋賀", "京都", "大阪", "兵庫", "奈良", "和歌山",
    "鳥取", "島根", "岡山", "広島", "山口", "徳島", "香川", "愛媛", "高知", "福岡",
    "佐賀", "長崎", "熊本", "大分", "宮崎", "鹿児島", "沖縄",
]
# 画面上の出荷者表記 → 集計用の名前。表記揺れはここに足していく。
SHIPPER_ALIASES = {
    "クラシック": "クラシック",
    "YMS": "YMS",
    "ワイエムエス": "YMS",
}

ITEM_START = re.compile(r"^([0-9]+)\t\s*$")
ORIGIN_LINE = re.compile(r"^(?P<origin>.+?)(?P<arrival>前入\d{1,2}:\d{2})?$")
LOTS = re.compile(r"(\d+)\s*口")
PRICE = re.compile(r"^¥([\d,]+)$")
TIER = re.compile(r"^(\d+)ケース\s*¥([\d,]+)$")


def nfkc(s):
    return unicodedata.normalize("NFKC", s)


def classify_origin(origin):
    if any(origin.startswith(p) for p in PREFECTURES) or origin in ("国産", "日本"):
        return "国産"
    if origin:
        return "外国産"
    return "要確認"


def normalize_shipper(raw):
    key = re.sub(r"(株式会社|株\)|\(株\)|有\)|\(有\))", "", nfkc(raw)).strip()
    for alias, name in SHIPPER_ALIASES.items():
        if alias.upper() in key.upper():
            return name
    return "その他"


def parse_header(text):
    m = re.search(r"売立日\s+(\d{2})月(\d{2})日", text)
    return f"{m.group(1)}/{m.group(2)}" if m else ""


def parse_block(lines):
    lines = [l for l in lines if l.strip()]
    ship_idx = next(
        i for i, l in enumerate(lines)
        if "\t" in l and not l.startswith("申し込む") and not l.startswith("¥")
    )
    head = [l.strip() for l in lines[: ship_idx - 1]]
    head = [l for l in head if not l.isdigit() and set(l) - set("■ ")]
    item = re.sub(r"^オススメ\s*", "", head[0]) if head else ""
    variety = head[1] if len(head) > 1 else ""
    note = " ".join(head[2:])  # 「夜間冷房対応」などの付記

    om = ORIGIN_LINE.match(lines[ship_idx - 1].strip())
    shipper_raw, grade = (lines[ship_idx].split("\t") + [""])[:2]

    rest = lines[ship_idx + 1:]
    packing = ""
    if rest and "\t" not in rest[0] and not rest[0].startswith("¥"):
        packing = rest.pop(0).strip()
    irisu = int(nfkc(rest.pop(0)).strip())

    price, tiers, lots = None, [], 0
    for l in rest:
        l = l.strip()
        if m := PRICE.match(l):
            price = int(m.group(1).replace(",", ""))
        elif m := TIER.match(l):
            tiers.append(f"{m.group(1)}c ¥{m.group(2)}")
        elif m := LOTS.search(l):
            lots = int(m.group(1))

    origin = om.group("origin")
    return {
        "品目": item,
        "品種": variety,
        "付記": note,
        "産地": origin,
        "国産/外国産": classify_origin(origin),
        "出荷者(表記)": shipper_raw.strip(),
        "出荷者": normalize_shipper(shipper_raw),
        "等階級": nfkc(grade).strip(),
        "荷姿": packing,
        "入数": irisu,
        "単価": price,
        "複数口単価": " / ".join(tiers),
        "申込可能口数": lots,
        "入荷": (om.group("arrival") or "").replace("前入", ""),
    }


def parse_text(text):
    date = parse_header(text)
    lines = text.splitlines()
    starts = [i for i, l in enumerate(lines) if ITEM_START.match(l)]
    rows = []
    for n, s in enumerate(starts):
        e = starts[n + 1] if n + 1 < len(starts) else len(lines)
        block = [l for l in lines[s + 1:e] if not l.startswith("No.\t")]
        row = parse_block(block)
        row = {"売立日": date, "No.": int(ITEM_START.match(lines[s]).group(1)), **row}
        rows.append(row)
    return rows


def write_workbook(rows, path):
    wb = Workbook()
    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    ws = wb.active
    ws.title = "明細"
    cols = list(rows[0].keys()) + ["本数"]
    ws.append(cols)
    irisu_col = get_column_letter(cols.index("入数") + 1)
    lots_col = get_column_letter(cols.index("申込可能口数") + 1)
    for r, row in enumerate(rows, start=2):
        ws.append(list(row.values()) + [f"={irisu_col}{r}*{lots_col}{r}"])
    for c in ws[1]:
        c.font, c.fill = bold, head_fill
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = base
    for i, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(8, len(name) * 2 + 2)
    ws.column_dimensions["E"].width = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    last = len(rows) + 1

    def rng(name):
        col = get_column_letter(cols.index(name) + 1)
        return f"明細!${col}$2:${col}${last}"

    s = wb.create_sheet("集計")
    s.column_dimensions["A"].width = 16
    for col in "BCDE":
        s.column_dimensions[col].width = 14

    def table(top, title, key, labels):
        s.cell(top, 1, title).font = bold
        hdr = [key, "行数", "口数", "本数", "本数シェア"]
        for j, h in enumerate(hdr, start=1):
            c = s.cell(top + 1, j, h)
            c.font, c.fill = bold, head_fill
        first = top + 2
        for i, label in enumerate(labels):
            r = first + i
            s.cell(r, 1, label)
            s.cell(r, 2, f'=COUNTIFS({rng(key)},$A{r})')
            s.cell(r, 3, f'=SUMIFS({rng("申込可能口数")},{rng(key)},$A{r})')
            s.cell(r, 4, f'=SUMIFS({rng("本数")},{rng(key)},$A{r})')
        tot = first + len(labels)
        s.cell(tot, 1, "合計").font = bold
        for j in (2, 3, 4):
            L = get_column_letter(j)
            s.cell(tot, j, f"=SUM({L}{first}:{L}{tot - 1})").font = bold
        for r in range(first, tot + 1):
            s.cell(r, 5, f"=IF($D${tot}=0,0,D{r}/$D${tot})").number_format = "0.0%"
            for j in (2, 3, 4):
                s.cell(r, j).number_format = "#,##0"
        for row in s.iter_rows(min_row=first, max_row=tot - 1):
            for c in row:
                c.font = base
        return tot + 2

    nxt = table(1, "国産／外国産", "国産/外国産", ["国産", "外国産", "要確認"])
    nxt = table(nxt, "出荷者別", "出荷者", ["クラシック", "YMS", "その他"])
    origins = sorted({r["産地"] for r in rows})
    nxt = table(nxt, "産地別", "産地", origins)
    s.cell(nxt, 1, "注: 本数 = 入数 × 申込可能口数。取得時点の残口数なので、上場総量ではない。").font = Font(
        name="Arial", size=9, italic=True
    )
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("-o", "--output", default="olive_offer.xlsx")
    a = ap.parse_args()
    rows = []
    for p in a.inputs:
        rows += parse_text(Path(p).read_text(encoding="utf-8"))
    if not rows:
        sys.exit("商品行が見つかりませんでした")
    write_workbook(rows, a.output)
    print(f"{len(rows)} 行 → {a.output}")


if __name__ == "__main__":
    main()
